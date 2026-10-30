"""
Planilha e curadoria dos deslistados dos EUA desde 2010 (docs/planos/PLANO_ATIVOS_INTERNACIONAIS.md, C6, C9 e D2).

    python carga/scripts/levantamento/internacional_deslistados.py planilha
        deslistados-previa.json (gravada pela tarefa FOREIGN_DELISTED) -> deslistados-revisao.xlsx
    python carga/scripts/levantamento/internacional_deslistados.py curadoria [planilha.xlsx]
        colunas amarelas "CURADORIA ..." -> curadoria-deslistados.json (a tarefa aplica na proxima rodada)

Abas: Resumo, Cobertura (ano x tipo: papeis de capital e quantos tem ticker), Deslistados, Sem ticker, Fora (notas,
preferenciais, warrants, troca de bolsa...), Categorias. A chave da curadoria e o numero de acesso do Form 25. Ao gerar
de novo, a curadoria ja feita (json e planilha anterior, copiada para .bak) volta para as colunas amarelas.
Precisa de openpyxl.
"""
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[3]
DIR = ROOT / "carga" / "data" / "levantamento" / "internacional"
PREVIA = DIR / "deslistados-previa.json"
PREVIA_VIVO = DIR / "catalogo-previa.json"
SAIDA = DIR / "deslistados-revisao.xlsx"
CURADORIA = DIR / "curadoria-deslistados.json"
TIPOS = ["STOCK", "REIT", "ETF_INTERNACIONAL"]
CUR = ["CURADORIA ticker", "CURADORIA tipo", "CURADORIA categoria", "CURADORIA excluir (S)", "CURADORIA obs"]
CHAVES = ["ticker", "tipo", "categoria", "excluir", "obs"]
ABAS_CURADAS = ["Deslistados", "Sem ticker", "Fora"]
AMARELO = PatternFill("solid", fgColor="FFF2CC")
CINZA = PatternFill("solid", fgColor="D9D9D9")
NEGRITO = Font(bold=True)


def texto(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def anterior():
    out = {}
    if CURADORIA.exists():
        for it in json.loads(CURADORIA.read_text(encoding="utf-8")).get("itens", []):
            out[it["acesso"]] = {k: it.get(k) for k in CHAVES}
    if SAIDA.exists():
        wb = load_workbook(SAIDA, read_only=True, data_only=True)
        for aba in ABAS_CURADAS:
            if aba not in wb.sheetnames:
                continue
            rows = wb[aba].iter_rows(values_only=True)
            cab = list(next(rows, []))
            if "acesso" not in cab:
                continue
            ix = {c: i for i, c in enumerate(cab)}
            for r in rows:
                vals = {k: r[ix[c]] for c, k in zip(CUR, CHAVES) if c in ix}
                if r[ix["acesso"]] and any(texto(v) for v in vals.values()):
                    out[r[ix["acesso"]]] = vals
    return out


def ajusta(ws, larguras):
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, cell in enumerate(ws[1], start=1):
        cell.font = NEGRITO
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = AMARELO if str(cell.value or "").startswith("CURADORIA") else CINZA
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(cell.value, 14)


def valida(ws, cab):
    n = max(ws.max_row, 2)
    ix = {c: i + 1 for i, c in enumerate(cab)}
    for col, lista in (("CURADORIA tipo", ",".join(TIPOS)), ("CURADORIA excluir (S)", "S")):
        dv = DataValidation(type="list", formula1=f'"{lista}"', allow_blank=True)
        ws.add_data_validation(dv)
        letra = get_column_letter(ix[col])
        dv.add(f"{letra}2:{letra}{n}")


def planilha():
    if not PREVIA.exists():
        sys.exit(f"Previa ausente: {PREVIA}. Rode a tarefa FOREIGN_DELISTED (WRITE=false) antes.")
    p = json.loads(PREVIA.read_text(encoding="utf-8"))
    cur = anterior()
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    linhas = [["Deslistados dos EUA - planilha de revisao (PLANO_ATIVOS_INTERNACIONAIS, C6/D2)"],
              ["Periodo", f"{p['de']}-{p['ate']}"], ["Previa gerada em", p["geradoEm"]],
              ["Gravou nesta rodada", "sim" if p["gravou"] else "nao (WRITE=false)"],
              ["Form 25 lidos", p["form25"]],
              ["Papeis de capital (nas bolsas da D1)", len(p["deslistados"])],
              ["Com ticker", sum(1 for r in p["deslistados"] if r.get("ticker"))],
              ["Curadoria aplicada", p["curadoria"]["aplicados"]], [],
              ["Fora, por motivo", "quantidade"]]
    linhas += [[k, v] for k, v in p["foraPorMotivo"].items()]
    linhas += [[], ["Como revisar"],
               ["1", "Abas Deslistados, Sem ticker e Fora: preencha so as colunas amarelas (CURADORIA ...)."],
               ["2", "ticker: o do papel na data da saida; tipo: STOCK, REIT ou ETF_INTERNACIONAL; categoria: aba Categorias."],
               ["3", "excluir: S tira o papel. Na aba Fora, preencher ticker e tipo inclui o papel."],
               ["4", "Rode: python carga/scripts/levantamento/internacional_deslistados.py curadoria"],
               ["5", "Depois a tarefa FOREIGN_DELISTED aplica (WRITE=false confere na previa; WRITE=true grava)."]]
    for l in linhas:
        ws.append(l)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 90
    ws["A1"].font = Font(bold=True, size=13)

    wc = wb.create_sheet("Cobertura")
    wc.append(["ano", "tipo", "papeis de capital", "com ticker", "cobertura %"])
    for c in p["cobertura"]:
        wc.append([c["ano"], c["tipo"], c["papeis"], c["comTicker"],
                   round(100.0 * c["comTicker"] / c["papeis"], 1) if c["papeis"] else None])
    ajusta(wc, {})

    cab = ["acesso", "ticker", "empresa", "classe", "bolsa", "saida", "arquivado", "tipo", "categoria",
           "origem ticker", "documento do ticker", "origem categoria", "SIC", "descricao SIC", "setor Nareit", "CIK",
           "acao", "obs"] + CUR
    larg = {"empresa": 34, "classe": 34, "documento do ticker": 40, "descricao SIC": 28, "obs": 30, "acesso": 22,
            "CURADORIA categoria": 24, "CURADORIA obs": 28}
    for aba, filtro in (("Deslistados", lambda r: r.get("ticker")), ("Sem ticker", lambda r: not r.get("ticker"))):
        w = wb.create_sheet(aba)
        w.append(cab)
        for r in (x for x in p["deslistados"] if filtro(x)):
            c = cur.get(r["accession"], {})
            w.append([r["accession"], r.get("ticker"), r["company"], r.get("securityClass"), r["exchange"],
                      r["delistedOn"], r["filed"], r["type"], r.get("category"), r.get("tickerSource"),
                      r.get("tickerDocument"), r.get("categorySource"), r.get("sic"), r.get("sicDescription"),
                      r.get("nareitSector"), r["cik"], r["action"], r.get("note"),
                      c.get("ticker"), c.get("tipo"), c.get("categoria"),
                      "S" if c.get("excluir") in (True, "S", "s") else None, c.get("obs")])
        ajusta(w, larg)
        valida(w, cab)

    cab_f = ["acesso", "empresa", "classe", "bolsa", "arquivado", "motivo", "CIK"] + CUR
    wf = wb.create_sheet("Fora")
    wf.append(cab_f)
    for o in p["fora"]:
        c = cur.get(o["accession"], {})
        wf.append([o["accession"], o["company"], o.get("securityClass"), o.get("exchange"), o["filed"], o["reason"],
                   o["cik"], c.get("ticker"), c.get("tipo"), c.get("categoria"),
                   "S" if c.get("excluir") in (True, "S", "s") else None, c.get("obs")])
    ajusta(wf, {"empresa": 34, "classe": 40, "motivo": 36, "acesso": 22})
    valida(wf, cab_f)

    wk = wb.create_sheet("Categorias")
    wk.append(["tipo", "categoria", "descricao"])
    if PREVIA_VIVO.exists():
        for c in json.loads(PREVIA_VIVO.read_text(encoding="utf-8"))["categorias"]:
            wk.append([c["type"], c["category"], c["description"]])
    ajusta(wk, {"categoria": 28, "descricao": 40})

    if SAIDA.exists():
        shutil.copy2(SAIDA, SAIDA.with_suffix(".xlsx.bak"))
    wb.save(SAIDA)
    print(f"Planilha: {SAIDA} ({len(p['deslistados'])} de capital, {len(p['fora'])} fora, curadoria previa {len(cur)})")


def curadoria():
    origem = Path(sys.argv[2]) if len(sys.argv) > 2 else SAIDA
    wb = load_workbook(origem, read_only=True, data_only=True)
    cats = {t: set() for t in TIPOS}
    for r in list(wb["Categorias"].iter_rows(values_only=True))[1:]:
        if r[0] in cats and r[1]:
            cats[r[0]].add(r[1])
    itens, erros = [], []
    for aba in ABAS_CURADAS:
        rows = wb[aba].iter_rows(values_only=True)
        cab = list(next(rows))
        ix = {c: i for i, c in enumerate(cab)}
        for n, r in enumerate(rows, start=2):
            v = {k: texto(r[ix[c]]) for c, k in zip(CUR, CHAVES)}
            if not any(v.values()):
                continue
            acc = r[ix["acesso"]]
            tipo = (v["tipo"] or "").upper() or None
            tipo_linha = r[ix["tipo"]] if "tipo" in ix else None
            if tipo and tipo not in TIPOS:
                erros.append(f"{aba} linha {n}: tipo desconhecido {tipo!r}")
            cat = (v["categoria"] or "").upper() or None
            alvo = tipo or tipo_linha
            if cat and (alvo not in cats or cat not in cats[alvo]):
                erros.append(f"{aba} linha {n}: categoria {cat} nao existe em {alvo}")
            excluir = (v["excluir"] or "").upper() in ("S", "SIM", "X")
            if aba == "Fora" and not excluir and not (v["ticker"] and tipo):
                erros.append(f"{aba} linha {n}: para incluir um papel da aba Fora, preencha ticker e tipo")
            itens.append({"acesso": acc, "aba": aba, "ticker": (v["ticker"] or "").upper() or None, "tipo": tipo,
                          "categoria": cat, "excluir": excluir, "obs": v["obs"]})
    if erros:
        print("\n".join(erros))
        sys.exit(f"{len(erros)} erro(s): {CURADORIA.name} nao foi gravado.")
    CURADORIA.write_text(json.dumps({"geradoEm": datetime.now().isoformat(timespec="seconds"), "planilha": origem.name,
                                     "itens": itens}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{CURADORIA}: {len(itens)} itens")


if __name__ == "__main__":
    {"planilha": planilha, "curadoria": curadoria}.get(sys.argv[1] if len(sys.argv) > 1 else "",
                                                       lambda: sys.exit(__doc__))()
