"""
Gera a planilha de revisao do catalogo do exterior (docs/planos/PLANO_ATIVOS_INTERNACIONAIS.md, C9 e D12) a partir da
previa gravada pela tarefa FOREIGN_CATALOG:

    carga/data/levantamento/internacional/catalogo-previa.json  ->  .../catalogo-revisao.xlsx

Uso (na raiz do repositorio; precisa de openpyxl):
    python carga/scripts/levantamento/internacional_planilha.py

Abas: Resumo, Stocks, REITs, ETFs, Deslistados, Categorias, Colisoes, Divergencias REIT, Excluidos, Logos a conferir.
As colunas amarelas "CURADORIA ..." sao as do Guilherme; depois de preencher, rode internacional_curadoria.py, que gera
curadoria.json (a tarefa aplica com a origem CURATION, nunca sobrescrita).

A curadoria nao se perde ao gerar de novo: o que ja esta em curadoria.json e o que estiver preenchido na planilha
anterior (copiada para .bak) volta para as colunas amarelas - a planilha anterior vence.
"""
import csv
import json
import shutil
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[3]
DIR = ROOT / "carga" / "data" / "levantamento" / "internacional"
PREVIA = DIR / "catalogo-previa.json"
SAIDA = DIR / "catalogo-revisao.xlsx"
CURADORIA = DIR / "curadoria.json"
LOGOS = ROOT / "docs" / "internacional" / "logos_relatorio.csv"

TIPOS = ["STOCK", "REIT", "ETF_INTERNACIONAL"]
ABA_TIPO = {"STOCK": "Stocks", "REIT": "REITs", "ETF_INTERNACIONAL": "ETFs"}
CUR = ["CURADORIA tipo", "CURADORIA categoria", "CURADORIA ticker", "CURADORIA excluir (S)", "CURADORIA obs"]
CUR_CHAVES = ["tipo", "categoria", "tickerNovo", "excluir", "obs"]

AMARELO = PatternFill("solid", fgColor="FFF2CC")
CINZA = PatternFill("solid", fgColor="D9D9D9")
VERMELHO = PatternFill("solid", fgColor="F8CBAD")
LARANJA = PatternFill("solid", fgColor="FCE4D6")
NEGRITO = Font(bold=True)


def curadoria_anterior():
    """ticker -> {tipo, categoria, tickerNovo, excluir, obs}: o json e, por cima, a planilha anterior."""
    out = {}
    if CURADORIA.exists():
        for it in json.loads(CURADORIA.read_text(encoding="utf-8")).get("itens", []):
            out[it["ticker"].upper()] = {k: it.get(k) for k in CUR_CHAVES}
    if SAIDA.exists():
        wb = load_workbook(SAIDA, read_only=True, data_only=True)
        for aba in list(ABA_TIPO.values()) + ["Excluidos"]:
            if aba not in wb.sheetnames:
                continue
            rows = wb[aba].iter_rows(values_only=True)
            cab = list(next(rows, []))
            if "ticker" not in cab:
                continue
            ix = {c: i for i, c in enumerate(cab)}
            for r in rows:
                t = r[ix["ticker"]]
                if not t:
                    continue
                vals = {k: r[ix[c]] for c, k in zip(CUR, CUR_CHAVES) if c in ix}
                if any(v not in (None, "") for v in vals.values()):
                    out[str(t).upper()] = vals
    return out


def ajusta(ws, larguras=None):
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, cell in enumerate(ws[1], start=1):
        cell.font = NEGRITO
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if str(cell.value or "").startswith("CURADORIA"):
            cell.fill = AMARELO
        else:
            cell.fill = CINZA
        w = (larguras or {}).get(cell.value, 14)
        ws.column_dimensions[get_column_letter(i)].width = w


def escreve(ws, cab, linhas):
    ws.append(cab)
    for l in linhas:
        ws.append(l)


def main():
    if not PREVIA.exists():
        sys.exit(f"Previa ausente: {PREVIA}. Rode a tarefa FOREIGN_CATALOG (WRITE=false) antes.")
    p = json.loads(PREVIA.read_text(encoding="utf-8"))
    cur = curadoria_anterior()
    cats = defaultdict(list)
    for c in p["categorias"]:
        cats[c["type"]].append(c)

    wb = Workbook()
    # ---------------------------------------------------------------- Resumo
    ws = wb.active
    ws.title = "Resumo"
    r = p["resumo"]
    linhas = [
        ["Catalogo do exterior - planilha de revisao (PLANO_ATIVOS_INTERNACIONAIS, C9/D12)"],
        ["Fotografia da lista da bolsa", p["fotografia"]],
        ["Previa gerada em", p["geradoEm"]],
        ["Gravou o catalogo nesta rodada", "sim" if p["gravou"] else "nao (WRITE=false)"],
        [],
        ["Stocks", r["STOCK"]], ["REITs", r["REIT"]], ["ETFs", r["ETF_INTERNACIONAL"]],
        ["Excluidos", r["excluidos"]], ["Sem categoria (aguardam a curadoria)", r["semCategoria"]],
        [],
        ["Acao prevista", "quantidade"],
    ]
    linhas += [[k, v] for k, v in r["acoes"].items()]
    linhas += [[], ["Tipo / origem da categoria", "quantidade"]]
    linhas += [[k, v] for k, v in r["origemCategoria"].items()]
    linhas += [[], ["Curadoria aplicada nesta rodada", p["curadoria"]["aplicados"]],
               ["Problemas da curadoria", len(p["curadoria"]["problemas"])]]
    linhas += [["", x] for x in p["curadoria"]["problemas"]]
    linhas += [[], ["Como revisar"],
               ["1", "Nas abas Stocks, REITs, ETFs e Excluidos, preencha so as colunas amarelas (CURADORIA ...)."],
               ["2", "tipo: STOCK, REIT ou ETF_INTERNACIONAL (no Excluido, inclui o papel com esse tipo)."],
               ["3", "categoria: um nome da aba Categorias do tipo da linha (ou do tipo curado)."],
               ["4", "ticker: so para corrigir o ticker gravado; excluir: S tira o papel do catalogo."],
               ["5", "Salve e rode: python carga/scripts/levantamento/internacional_curadoria.py"],
               ["6", "Depois a tarefa FOREIGN_CATALOG aplica (WRITE=false para conferir na previa; WRITE=true grava)."]]
    for l in linhas:
        ws.append(l)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 90
    ws["A1"].font = Font(bold=True, size=13)

    # ---------------------------------------------------------------- Categorias
    wc = wb.create_sheet("Categorias")
    contagem = Counter((a["type"], a["category"]) for a in p["ativos"])
    escreve(wc, ["tipo", "categoria", "descricao", "papeis"],
            [[c["type"], c["category"], c["description"], contagem.get((c["type"], c["category"]), 0)]
             for c in p["categorias"]])
    ajusta(wc, {"descricao": 40, "categoria": 30})
    faixa = {}
    for t in TIPOS:
        linhas_t = [i + 2 for i, c in enumerate(p["categorias"]) if c["type"] == t]
        if linhas_t:
            faixa[t] = f"Categorias!$B${min(linhas_t)}:$B${max(linhas_t)}"
    # BDR x SIC: onde a tabela SIC discorda do BDR da mesma empresa (para revisar a tabela)
    dif = defaultdict(list)
    for a in p["ativos"]:
        if a["type"] == "STOCK" and a.get("bdrCategory") and a.get("sicCategory") \
                and a["bdrCategory"] != a["sicCategory"]:
            dif[(a.get("sic"), a.get("sicDescription"), a["sicCategory"], a["bdrCategory"])].append(a["ticker"])
    col = 7
    wc.cell(row=1, column=col, value="BDR x SIC (empresas com BDR em que a tabela SIC daria outro setor)").font = NEGRITO
    cab = ["SIC", "descricao SIC", "setor pela tabela SIC", "setor do BDR", "empresas", "tickers"]
    for j, h in enumerate(cab):
        wc.cell(row=2, column=col + j, value=h).font = NEGRITO
    for i, (k, v) in enumerate(sorted(dif.items(), key=lambda kv: -len(kv[1])), start=3):
        for j, val in enumerate(list(k) + [len(v), ", ".join(v[:12])]):
            wc.cell(row=i, column=col + j, value=val)
    for j, w in enumerate([8, 40, 28, 28, 10, 50]):
        wc.column_dimensions[get_column_letter(col + j)].width = w

    # ---------------------------------------------------------------- Stocks, REITs, ETFs
    base = ["ticker", "nome", "bolsa", "acao", "categoria", "origem categoria"]
    extra = {
        "STOCK": ["categoria BDR", "BDR", "categoria SIC", "SIC", "descricao SIC", "setor Nareit", "tipo OpenFIGI",
                  "origem tipo", "CIK"],
        "REIT": ["setor Nareit", "categoria SIC", "SIC", "descricao SIC", "tipo OpenFIGI", "origem tipo", "CIK"],
        "ETF_INTERNACIONAL": ["tipo OpenFIGI", "origem tipo", "CIK", "serie SEC"],
    }
    campo = {"categoria BDR": "bdrCategory", "BDR": "bdrTicker", "categoria SIC": "sicCategory", "SIC": "sic",
             "descricao SIC": "sicDescription", "setor Nareit": "nareitSector", "tipo OpenFIGI": "figiSecurityType",
             "origem tipo": "typeSource", "CIK": "cik", "serie SEC": "seriesId"}
    for t in TIPOS:
        ws = wb.create_sheet(ABA_TIPO[t])
        cab = base + extra[t] + ["FIGI", "ticker anterior", "obs"] + CUR
        ws.append(cab)
        for a in sorted((a for a in p["ativos"] if a["type"] == t), key=lambda a: a["ticker"]):
            c = cur.get(a["ticker"].upper(), {})
            linha = [a["ticker"], a["name"], a["exchange"], a["action"], a.get("category"), a.get("categorySource")]
            linha += [a.get(campo[x]) for x in extra[t]]
            linha += [a.get("compositeFigi"), a.get("previousTicker"), a.get("note")]
            linha += [c.get("tipo"), c.get("categoria"), c.get("tickerNovo"),
                      "S" if c.get("excluir") in (True, "S", "s") else None, c.get("obs")]
            ws.append(linha)
            row = ws.max_row
            if not a.get("category"):
                for cell in ws[row][:6]:
                    cell.fill = VERMELHO
            elif t == "STOCK" and a.get("bdrCategory") and a.get("sicCategory") \
                    and a["bdrCategory"] != a["sicCategory"]:
                ws.cell(row=row, column=5).fill = LARANJA
        ajusta(ws, {"nome": 45, "obs": 40, "descricao SIC": 30, "setor Nareit": 28, "categoria": 24,
                    "categoria BDR": 22, "categoria SIC": 22, "CURADORIA categoria": 24, "CURADORIA obs": 30})
        valida(ws, cab, faixa.get(t))

    # ---------------------------------------------------------------- Deslistados
    wd = wb.create_sheet("Deslistados")
    escreve(wd, ["ticker", "tipo", "nome", "situacao", "origem", "saida", "visto por ultimo", "prova", "id do ativo"],
            [[g["ticker"], g["type"], g["name"], g["status"], g.get("source"), g.get("delistedOn"), g.get("lastSeen"),
              g.get("proof"), g["assetId"]] for g in p["deslistados"]])
    ajusta(wd, {"nome": 40, "prova": 90})

    # ---------------------------------------------------------------- Colisoes
    wl = wb.create_sheet("Colisoes")
    escreve(wl, ["ticker", "tipo", "outro tipo com o mesmo ticker", "nome do outro"],
            [[c["ticker"], c["type"], c["otherType"], c["otherName"]] for c in p["colisoes"]])
    ajusta(wl, {"nome do outro": 40, "outro tipo com o mesmo ticker": 30})

    # ---------------------------------------------------------------- Divergencias REIT
    wr = wb.create_sheet("Divergencias REIT")
    escreve(wr, ["ticker", "nome", "tipo (OpenFIGI decide)", "tipo OpenFIGI", "setor Nareit", "SIC", "obs"],
            [[d["ticker"], d["name"], d["type"], d.get("figiSecurityType"), d.get("nareitSector"), d.get("sic"),
              d["note"]] for d in p["divergenciasReit"]])
    ajusta(wr, {"nome": 40, "obs": 45, "setor Nareit": 30})

    # ---------------------------------------------------------------- Excluidos
    we = wb.create_sheet("Excluidos")
    cab = ["ticker", "simbolo", "nome", "bolsa", "motivo", "tipo OpenFIGI"] + CUR
    we.append(cab)
    for x in sorted(p["excluidos"], key=lambda x: (x["reason"], x["ticker"])):
        c = cur.get(x["ticker"].upper(), {})
        we.append([x["ticker"], x["symbol"], x["name"], x["exchange"], x["reason"], x.get("figiSecurityType"),
                   c.get("tipo"), c.get("categoria"), c.get("tickerNovo"),
                   "S" if c.get("excluir") in (True, "S", "s") else None, c.get("obs")])
    ajusta(we, {"nome": 60, "motivo": 28, "CURADORIA categoria": 24})
    valida(we, cab, None)

    # ---------------------------------------------------------------- Logos a conferir
    wg = wb.create_sheet("Logos a conferir")
    linhas = []
    if LOGOS.exists():
        with open(LOGOS, encoding="utf-8") as f:
            for row in csv.DictReader(f):
                if "conferir" in (row.get("acao") or ""):
                    linhas.append([row["ticker"], row["tipo"], row.get("etapa"), row["acao"], row.get("origem")])
    escreve(wg, ["ticker", "tipo", "etapa", "acao", "origem"], linhas)
    ajusta(wg, {"acao": 30, "origem": 60})

    if SAIDA.exists():
        shutil.copy2(SAIDA, SAIDA.with_suffix(".xlsx.bak"))
    DIR.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    print(f"Planilha: {SAIDA}")
    print(f"  Stocks {r['STOCK']}, REITs {r['REIT']}, ETFs {r['ETF_INTERNACIONAL']}, excluidos {r['excluidos']}, "
          f"sem categoria {r['semCategoria']}, deslistados {len(p['deslistados'])}, curadoria previa {len(cur)}")


def valida(ws, cab, faixa_categoria):
    """Listas nas colunas da curadoria: tipo, categoria (do tipo da aba) e excluir."""
    n = max(ws.max_row, 2)
    ix = {c: i + 1 for i, c in enumerate(cab)}
    tipo = DataValidation(type="list", formula1='"' + ",".join(TIPOS) + '"', allow_blank=True)
    ws.add_data_validation(tipo)
    letra = get_column_letter(ix["CURADORIA tipo"])
    tipo.add(f"{letra}2:{letra}{n}")
    exc = DataValidation(type="list", formula1='"S"', allow_blank=True)
    ws.add_data_validation(exc)
    letra = get_column_letter(ix["CURADORIA excluir (S)"])
    exc.add(f"{letra}2:{letra}{n}")
    if faixa_categoria:
        cat = DataValidation(type="list", formula1=f"={faixa_categoria}", allow_blank=True)
        cat.showErrorMessage = False
        ws.add_data_validation(cat)
        letra = get_column_letter(ix["CURADORIA categoria"])
        cat.add(f"{letra}2:{letra}{n}")


if __name__ == "__main__":
    main()
