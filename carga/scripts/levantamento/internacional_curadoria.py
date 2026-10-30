"""
Converte as colunas "CURADORIA ..." da planilha de revisao do exterior em curadoria.json
(docs/planos/PLANO_ATIVOS_INTERNACIONAIS.md, C9):

    carga/data/levantamento/internacional/catalogo-revisao.xlsx  ->  .../curadoria.json

Uso (na raiz do repositorio; precisa de openpyxl):
    python carga/scripts/levantamento/internacional_curadoria.py
    python carga/scripts/levantamento/internacional_curadoria.py outra_planilha.xlsx

Le as abas Stocks, REITs, ETFs e Excluidos. Uma linha entra quando alguma coluna amarela esta preenchida. Confere o
tipo (STOCK, REIT, ETF_INTERNACIONAL) e a categoria (a da aba Categorias do tipo da linha ou do tipo curado); com erro,
nada e gravado. A tarefa FOREIGN_CATALOG le o arquivo na proxima rodada e aplica com a origem CURATION.
Depois de gerar: docker compose --profile carga run --rm data push
"""
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
DIR = ROOT / "carga" / "data" / "levantamento" / "internacional"
PLANILHA = DIR / "catalogo-revisao.xlsx"
SAIDA = DIR / "curadoria.json"

TIPOS = {"STOCK", "REIT", "ETF_INTERNACIONAL"}
ABAS = {"Stocks": "STOCK", "REITs": "REIT", "ETFs": "ETF_INTERNACIONAL", "Excluidos": None}
COLUNAS = {"CURADORIA tipo": "tipo", "CURADORIA categoria": "categoria", "CURADORIA ticker": "tickerNovo",
           "CURADORIA excluir (S)": "excluir", "CURADORIA obs": "obs"}


def texto(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def main():
    origem = Path(sys.argv[1]) if len(sys.argv) > 1 else PLANILHA
    if not origem.exists():
        sys.exit(f"Planilha ausente: {origem}")
    wb = load_workbook(origem, read_only=True, data_only=True)
    categorias = {t: set() for t in TIPOS}
    for r in list(wb["Categorias"].iter_rows(values_only=True))[1:]:
        if r[0] in TIPOS and r[1]:
            categorias[r[0]].add(r[1])
    itens, erros, vistos = [], [], set()
    for aba, tipo_aba in ABAS.items():
        if aba not in wb.sheetnames:
            continue
        rows = wb[aba].iter_rows(values_only=True)
        cab = list(next(rows))
        ix = {c: i for i, c in enumerate(cab)}
        for n, r in enumerate(rows, start=2):
            ticker = texto(r[ix["ticker"]])
            if not ticker:
                continue
            v = {k: texto(r[ix[c]]) for c, k in COLUNAS.items() if c in ix}
            if not any(v.values()):
                continue
            ticker = ticker.upper()
            if ticker in vistos:
                erros.append(f"{aba} linha {n}: {ticker} repetido em outra aba")
                continue
            vistos.add(ticker)
            tipo = v["tipo"].upper() if v.get("tipo") else None
            if tipo and tipo not in TIPOS:
                erros.append(f"{aba} linha {n}: tipo desconhecido {v['tipo']!r}")
            alvo = tipo or tipo_aba
            cat = v["categoria"].upper() if v.get("categoria") else None
            if cat:
                if alvo is None:
                    erros.append(f"{aba} linha {n}: {ticker} tem categoria sem tipo (preencha CURADORIA tipo)")
                elif cat not in categorias[alvo]:
                    erros.append(f"{aba} linha {n}: categoria {cat} nao existe em {alvo}")
            excluir = (v.get("excluir") or "").upper() in ("S", "SIM", "X", "TRUE")
            if v.get("excluir") and not excluir:
                erros.append(f"{aba} linha {n}: excluir deve ser S ({v['excluir']!r})")
            if aba == "Excluidos" and not tipo and not excluir:
                erros.append(f"{aba} linha {n}: {ticker} excluido pela regra so entra com CURADORIA tipo")
            itens.append({"ticker": ticker, "aba": aba, "tipo": tipo, "categoria": cat,
                          "tickerNovo": v["tickerNovo"].upper() if v.get("tickerNovo") else None,
                          "excluir": excluir, "obs": v.get("obs")})
    if erros:
        print("\n".join(erros))
        sys.exit(f"{len(erros)} erro(s): {SAIDA.name} nao foi gravado.")
    SAIDA.write_text(json.dumps({"geradoEm": datetime.now().isoformat(timespec="seconds"),
                                 "planilha": origem.name, "itens": itens}, ensure_ascii=False, indent=2),
                     encoding="utf-8")
    print(f"{SAIDA}: {len(itens)} itens ({sum(1 for i in itens if i['excluir'])} excluidos, "
          f"{sum(1 for i in itens if i['tipo'])} com tipo, {sum(1 for i in itens if i['categoria'])} com categoria, "
          f"{sum(1 for i in itens if i['tickerNovo'])} com ticker)")


if __name__ == "__main__":
    main()
