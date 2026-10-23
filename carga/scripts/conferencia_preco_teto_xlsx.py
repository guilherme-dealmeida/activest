"""Monta a planilha de conferencia do Preco Teto (docs/planos/PLANO_PRECO_TETO.md, 8.2).

Le o JSON gerado por conferencia_preco_teto.mjs e escreve um .xlsx no LAYOUT da PRECO TETO V2 (aba "Valuation V2",
colunas A..AA com as MESMAS formulas da planilha), uma linha por calculo, e ao lado as colunas com o resultado do
Activest e a diferenca. Ao abrir no Excel as formulas recalculam: a diferenca tem de ser so de arredondamento.

TX CRES (E) entra como formula =D*(1-C) — a convencao do "Leia-me" da V2 —, igual ao Activest.

Uso: python3 carga/scripts/conferencia_preco_teto_xlsx.py docs/preco-teto/conferencia_preco_teto_AAAA-MM-DD.json
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEAD = ['TICKER', 'NRO PAPEIS', 'PAYOUT', 'ROE', 'TX CRES', 'LL T-3', 'LL T-2', 'LL T-1', 'LL T',
        'LUCRO LÍQUIDO TTM', 'LUCRO PROJETADO / FC1', 'DESCONTO', 'MARGEM SEGURANÇA', 'VPL1', 'FCA2', 'VPL2',
        'FCA3', 'VPL3', 'G PERPÉTUO', 'FCAP', 'VPLP', 'MARKET CAP', 'PREÇO ATUAL', 'PREÇO TETO', 'MARGEM TETO',
        'PREÇO IDEAL', 'MARGEM IDEAL']
PV = ['PV LL TTM', 'PV VALOR PATRIMÔNIO', 'PV PREÇO TETO', 'PV PREÇO IDEAL', 'PV MARGEM TETO', 'PV MARGEM IDEAL',
      'DIF TETO (R$)', 'DIF IDEAL (R$)', 'CONFERE?', 'ORIGEM']

# formulas da V2 (aba Valuation V2, linha 2), com {r} = linha
FORMULAS = {
    'E': '=D{r}*(1-C{r})',
    'J': '=IF(COUNT(F{r}:I{r})=4,SUM(F{r}:I{r}),"")',
    'K': '=IF(J{r}<>"",J{r},"")',
    'N': '=IF(K{r}<>"",K{r}/(1+L{r}),"")',
    'O': '=IF(K{r}<>"",K{r}*(1+E{r}),"")',
    'P': '=IF(O{r}<>"",O{r}/(1+L{r})^2,"")',
    'Q': '=IF(O{r}<>"",O{r}*(1+E{r}),"")',
    'R': '=IF(Q{r}<>"",Q{r}/(1+L{r})^3,"")',
    'T': '=IF(Q{r}<>"",Q{r}*(1+S{r})/(L{r}-S{r}),"")',
    'U': '=IF(T{r}<>"",T{r}/(1+L{r})^3,"")',
    'V': '=IF(COUNT(N{r},P{r},R{r},U{r})=4,SUM(N{r},P{r},R{r},U{r}),"")',
    'X': '=IF(AND(V{r}<>"",B{r}>0),V{r}/B{r},"")',
    'Y': '=IF(AND(X{r}<>"",W{r}>0),X{r}/W{r}-1,"")',
    'Z': '=IF(AND(X{r}<>"",M{r}>=0),X{r}*(1-M{r}),"")',
    'AA': '=IF(AND(Z{r}<>"",W{r}>0),Z{r}/W{r}-1,"")',
}


def pct(v):
    return None if v is None else v / 100


def main(path):
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Valuation V2'
    yellow = PatternFill('solid', fgColor='FFF2CC')
    blue = PatternFill('solid', fgColor='DDEBF7')
    for c, h in enumerate(HEAD + PV, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        if c > len(HEAD):
            cell.fill = blue
    for i, linha in enumerate(data['linhas'], start=2):
        e, res = linha['entradas'], linha['resultado']
        q = [x.get('originalValue') for x in e['quarters']]
        valores = {
            'A': e['ticker'], 'B': e['shares'], 'C': pct(e['payoutPct']), 'D': pct(e['roePct']),
            'F': q[0], 'G': q[1], 'H': q[2], 'I': q[3], 'L': pct(e['discountRatePct']),
            'M': pct(e['safetyMarginPct']), 'S': pct(e['perpetualGrowthPct']), 'W': e['price'],
        }
        for col, v in valores.items():
            cell = ws[f'{col}{i}']
            cell.value = v
            cell.fill = yellow
        for col, f in FORMULAS.items():
            ws[f'{col}{i}'] = f.format(r=i)
        base = len(HEAD)
        pv = [res.get('netIncomeTtm'), res.get('equityValue'), res.get('ceilingPrice'), res.get('idealPrice'),
              pct(res.get('ceilingMarginPct')), pct(res.get('idealMarginPct'))]
        for k, v in enumerate(pv, start=1):
            ws.cell(row=i, column=base + k, value=v).fill = blue
        teto = get_column_letter(base + 3)
        ideal = get_column_letter(base + 4)
        ws.cell(row=i, column=base + 7, value=f'=IF(AND(ISNUMBER(X{i}),ISNUMBER({teto}{i})),{teto}{i}-X{i},"")')
        ws.cell(row=i, column=base + 8, value=f'=IF(AND(ISNUMBER(Z{i}),ISNUMBER({ideal}{i})),{ideal}{i}-Z{i},"")')
        dif = get_column_letter(base + 7)
        ws.cell(row=i, column=base + 9, value=f'=IF(ISNUMBER({dif}{i}),IF(ABS({dif}{i})<0.0001,"SIM","NÃO"),"—")')
        ws.cell(row=i, column=base + 10, value=linha['rotulo'])
        for col in ('C', 'D', 'E', 'L', 'M', 'S', 'Y', 'AA', get_column_letter(base + 5), get_column_letter(base + 6)):
            ws[f'{col}{i}'].number_format = '0.00%'
        for col in ('W', 'X', 'Z', teto, ideal, dif, get_column_letter(base + 8)):
            ws[f'{col}{i}'].number_format = '#,##0.0000'
        for col in ('B', 'F', 'G', 'H', 'I', 'J', 'K', 'N', 'O', 'P', 'Q', 'R', 'T', 'U', 'V',
                    get_column_letter(base + 1), get_column_letter(base + 2)):
            ws[f'{col}{i}'].number_format = '#,##0'
    for c in range(1, len(HEAD) + len(PV) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.column_dimensions[get_column_letter(len(HEAD) + len(PV))].width = 34
    ws.freeze_panes = 'B2'

    leia = wb.create_sheet('Leia-me')
    texto = [
        'CONFERÊNCIA PREÇO TETO — Activest × PREÇO TETO V2',
        f"Gerado em {data['geradoEm']} a partir de {data['api']}.",
        'Colunas A..AA: layout e fórmulas da aba Valuation V2 (entradas em amarelo).',
        'TX CRES (E) = ROE × (1 − payout), como sugere o Leia-me da V2.',
        'Colunas em azul: o resultado calculado pelo Activest com as mesmas entradas.',
        'DIF TETO/IDEAL = Activest − planilha. CONFERE? = SIM quando a diferença do teto é menor que R$ 0,0001.',
        'Linhas "(padrões do Activest)": entradas padrão do sistema (CVM, Selic meta, 3%, 15%, último fechamento).',
        'Linha "VALE3 (linha da planilha V2)": as entradas da planilha original.',
    ]
    for k, t in enumerate(texto, start=1):
        leia.cell(row=k, column=1, value=t)
    leia.column_dimensions['A'].width = 120
    out = Path(path).with_suffix('.xlsx')
    wb.save(out)
    print(f'Gravado: {out}')


if __name__ == '__main__':
    main(sys.argv[1])
