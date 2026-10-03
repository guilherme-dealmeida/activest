"""
Gera data/levantamento/datas-pagamento.json a partir da planilha do levantamento
(PLANO_CARGA_INICIAL.md, secao 11 - L1 a L4).

Uso (na raiz do repositorio; precisa de openpyxl):
    python scripts/levantamento/xlsx_para_json.py
    python scripts/levantamento/xlsx_para_json.py data/levantamento/origem/outra_planilha.xlsx

Chave natural de cada linha: ticker (+ ticker_fracionario), tipo, data_com, valor_por_acao e parcela.
O id_evento da planilha NAO vai para o arquivo: so existe no banco em que a planilha foi gerada.
Depois de gerar: docker compose --profile carga run --rm data push
"""
import json
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
ORIGEM = ROOT / "data" / "levantamento" / "origem"
SAIDA = ROOT / "data" / "levantamento" / "datas-pagamento.json"
STATUS = {"CONFIRMADO": "CONFIRMADO", "ACEITE NA ÚLTIMA REVISÃO": "ACEITE_ULTIMA_REVISAO", "A DEFINIR": "A_DEFINIR"}
TIPOS = {"DIVIDENDO", "JCP", "RENDIMENTO", "RESTITUICAO_CAPITAL", "AMORTIZACAO"}


def planilha():
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    xs = sorted(ORIGEM.glob("*.xlsx"))
    if len(xs) != 1:
        sys.exit(f"Esperava uma planilha em {ORIGEM}, achei {len(xs)}. Informe o arquivo.")
    return xs[0]


def dia(v, campo, linha):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    sys.exit(f"Linha {linha}: {campo} invalido: {v!r}")


def texto(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def main():
    origem = planilha()
    ws = load_workbook(origem, read_only=True, data_only=True)["Proventos"]
    rows = ws.iter_rows(values_only=True)
    cab = [c for c in next(rows)]
    ix = {c: i for i, c in enumerate(cab)}
    for c in ("ticker", "ticker_fracionario", "isin", "tipo", "data_com", "valor_por_acao", "parcela",
              "DATA_PAGAMENTO", "FONTE_PREENCHIMENTO", "STATUS", "STATUS GERAL", "REGRA_APLICADA"):
        if c not in ix:
            sys.exit(f"Coluna ausente na planilha: {c}")
    saida, chaves, erros = [], set(), []
    for n, r in enumerate(rows, start=2):
        g = lambda c: r[ix[c]]
        if g("ticker") is None:
            continue
        st = STATUS.get(texto(g("STATUS GERAL")))
        if st is None:
            erros.append(f"linha {n}: STATUS GERAL desconhecido {g('STATUS GERAL')!r}")
            continue
        tipo = texto(g("tipo"))
        if tipo not in TIPOS:
            erros.append(f"linha {n}: tipo desconhecido {tipo!r}")
        dc = dia(g("data_com"), "data_com", n)
        dp = dia(g("DATA_PAGAMENTO"), "DATA_PAGAMENTO", n)
        if st == "A_DEFINIR" and dp is not None:
            erros.append(f"linha {n}: A DEFINIR com data de pagamento")
        if st != "A_DEFINIR" and dp is None:
            erros.append(f"linha {n}: {st} sem data de pagamento")
        if dp is not None and dp < dc:
            erros.append(f"linha {n}: pagamento {dp} antes da data-com {dc}")
        valor = format(Decimal(str(g("valor_por_acao"))).quantize(Decimal("0.00000000001")), "f")
        parcela = int(g("parcela") or 1)
        chave = (texto(g("ticker")), tipo, dc, valor, parcela)
        if chave in chaves:
            erros.append(f"linha {n}: chave repetida {chave}")
        chaves.add(chave)
        saida.append({
            "ticker": texto(g("ticker")),
            "ticker_fracionario": texto(g("ticker_fracionario")),
            "isin": texto(g("isin")),
            "tipo": tipo,
            "data_com": dc.isoformat(),
            "valor_por_acao": valor,
            "parcela": parcela,
            "data_pagamento": dp.isoformat() if dp else None,
            "status": st,
            "fonte": texto(g("FONTE_PREENCHIMENTO")),
            "regra": " | ".join(x for x in (texto(g("STATUS")), texto(g("REGRA_APLICADA"))) if x),
        })
    if erros:
        print("\n".join(erros[:50]))
        sys.exit(f"{len(erros)} erro(s); nada foi gravado.")
    saida.sort(key=lambda x: (x["ticker"], x["data_com"], x["tipo"], x["parcela"], x["valor_por_acao"]))
    doc = {"origem": origem.name, "descricao": "Datas de pagamento levantadas para proventos de acoes sem data "
           "(PLANO_CARGA_INICIAL.md secao 11). status A_DEFINIR = sem data conhecida.", "linhas": saida}
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text(json.dumps(doc, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    cont = {}
    for x in saida:
        cont[x["status"]] = cont.get(x["status"], 0) + 1
    print(f"{len(saida)} linhas -> {SAIDA.relative_to(ROOT)}  {cont}")


if __name__ == "__main__":
    main()
